"""MOM Excel generator — one meeting, one workbook, exact VCP-format styling.

Replicates the styling of VCP-NEW-FORMAT/MOM VCP Meetings.xlsx precisely:
Aptos Narrow, steel-blue title bar (theme 3, white bold 16), light-yellow section
headers (bold 13.5), bold decision lines with spacer rows, bordered action table
with light-blue header row.

Usage:
  python momo_mom.py --json mom.json --out "MOM - Topic - 2026-09-01.xlsx"

mom.json:
{
  "title": "VCP Format Review",
  "date": "01-09-2026",
  "meeting_no": null,                # optional; omitted -> title only
  "decisions": ["...", "..."],
  "actions": [{"action": "...", "owner": "...", "timing": "..."}]
}
"""
import argparse
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

FONT = "Aptos Narrow"
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITLE_FILL = PatternFill("solid", fgColor=Color(theme=3, tint=0.2499771111178931))
TITLE_FONT = Font(name=FONT, size=16, bold=True, color=Color(theme=0))
SECTION_FILL = PatternFill("solid", fgColor="FFFFFF93")
SECTION_FONT = Font(name=FONT, size=13.5, bold=True)
HDR_FILL = PatternFill("solid", fgColor=Color(theme=3, tint=0.8999908444471572))
HDR_FONT = Font(name=FONT, size=11, bold=True)
BODY_FONT = Font(name=FONT, size=11)
DEC_FONT = Font(name=FONT, size=11, bold=True)

WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
WRAP_L = Alignment(wrap_text=True, vertical="center", horizontal="left")
LEFT = Alignment(horizontal="left", vertical="center")


def build(data, out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (data.get("title") or "MOM")[:28]

    for col, w in (("A", 5.5), ("B", 61.5), ("C", 41.2), ("D", 26.5)):
        ws.column_dimensions[col].width = w

    title = data.get("title", "Meeting")
    if data.get("meeting_no"):
        title = f"Meeting {data['meeting_no']} - {title}"
    if data.get("date"):
        title = f"{title} ({data['date']})"
    c = ws.cell(1, 1, title)
    c.font = TITLE_FONT
    c.alignment = LEFT
    for col in range(1, 5):
        ws.cell(1, col).fill = TITLE_FILL
    ws.row_dimensions[1].height = 21.0

    c = ws.cell(3, 1, "Key Decisions / Outcomes")
    c.font = SECTION_FONT
    c.alignment = LEFT
    for col in range(1, 5):
        ws.cell(3, col).fill = SECTION_FILL
    ws.row_dimensions[3].height = 17.5

    r = 5
    for i, d in enumerate(data.get("decisions", []), 1):
        n = ws.cell(r, 1, i)
        n.font = BODY_FONT
        n.alignment = WRAP_C
        t = ws.cell(r, 2, d)
        t.font = DEC_FONT
        t.alignment = WRAP_L
        ws.row_dimensions[r].height = max(15.0, 13.0 * (len(d) // 62 + 1))
        r += 2

    c = ws.cell(r, 1, "Action Items")
    c.font = SECTION_FONT
    c.alignment = LEFT
    for col in range(1, 5):
        ws.cell(r, col).fill = SECTION_FILL
    ws.row_dimensions[r].height = 17.5
    r += 2

    for ci, h in enumerate(["S. No.", "Action", "Owner", "Due / Timing"], 1):
        c = ws.cell(r, ci, h)
        c.font = HDR_FONT
        c.alignment = WRAP_C
        c.fill = HDR_FILL
        c.border = BORDER
    ws.row_dimensions[r].height = 29.0
    r += 1

    for i, a in enumerate(data.get("actions", []), 1):
        cells = [
            (1, i, WRAP_C),
            (2, a.get("action", ""), WRAP_L),
            (3, a.get("owner", "(to confirm)"), WRAP_L),
            (4, a.get("timing", "(to confirm)"), WRAP_L),
        ]
        for col, val, align in cells:
            c = ws.cell(r, col, val)
            c.font = BODY_FONT
            c.alignment = align
            c.border = BORDER
        ws.row_dimensions[r].height = max(29.0, 14.5 * (len(a.get("action", "")) // 62 + 1))
        r += 1

    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    wb.save(out)
    print("MOM WORKBOOK READY:", out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    with open(args.json, encoding="utf-8") as f:
        build(json.load(f), args.out)
