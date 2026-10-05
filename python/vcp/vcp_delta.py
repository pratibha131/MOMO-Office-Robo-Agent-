# -*- coding: utf-8 -*-
"""
Month-on-month delta for the VCP Presentation Generator.

compare(previous_model, current_model)  -> delta dict (themes, changes, summaries)
add_delta_slides(prs, view, s2, s4, ...) -> two slides after the themes slide:
    1. "Key theme movement vs previous tracker": cards + a native (editable) PowerPoint bar chart,
       previous vs current progress per key theme, in the template's colours.
    2. "What changed since the previous tracker": table of every action / KPI / dependency change.
summary_lines(delta) -> validation-summary section.

Matching rule (documented, deterministic): items are paired by normalised name (letters and digits
only, case-insensitive); exact matches first, then the best fuzzy match with a similarity of at least
0.85 for themes / KPIs / dependencies and 0.80 for actions. Everything unmatched is reported as new
or removed - never silently dropped.
"""
import difflib
import re
from pathlib import Path

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt

import vcp_deck as vd

TONE_PALETTE = {"up": "Completed", "down": "Delayed", "new": "In progress", "gone": "No data", "same": "No data"}
PREV_COLOUR = RGBColor(0xC9, 0xCF, 0xD6)
GRID_COLOUR = RGBColor(0xD9, 0xDE, 0xE5)
LABEL_MAX = 48


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def _match(prev_items, cur_items, key, threshold=0.85):
    """Pair items across the two trackers. Returns [(prev_or_None, cur_or_None), ...] in current order."""
    prev_left, cur_left, pairs = list(prev_items), list(cur_items), []
    for c in list(cur_left):
        for p in prev_left:
            if _norm(key(p)) and _norm(key(p)) == _norm(key(c)):
                pairs.append((p, c))
                prev_left.remove(p)
                cur_left.remove(c)
                break
    cands = []
    for c in cur_left:
        for p in prev_left:
            r = difflib.SequenceMatcher(None, _norm(key(p)), _norm(key(c))).ratio()
            if r >= threshold:
                cands.append((r, p, c))
    cands.sort(key=lambda x: -x[0])
    used_p, used_c = set(), set()
    for r, p, c in cands:
        if id(p) in used_p or id(c) in used_c:
            continue
        pairs.append((p, c))
        used_p.add(id(p))
        used_c.add(id(c))
    for c in cur_left:
        if id(c) not in used_c:
            pairs.append((None, c))
    for p in prev_left:
        if id(p) not in used_p:
            pairs.append((p, None))
    return pairs


def theme_pct(t, none_if_empty=False):
    tp = next((r["theme_progress"] for r in t["rows"] if r["theme_progress"] is not None), None)
    if tp is not None:
        return round(100 * tp)
    vals = [r["progress"] for r in t["rows"] if r["progress"] is not None]
    if not vals and none_if_empty:
        return None
    return vd.pct([r["progress"] for r in t["rows"]])


def _p(v):
    return None if v is None else round(100 * v)


def _row(year, theme, change, movement, tone):
    return {"year": str(year), "theme": theme or "", "change": change, "movement": movement, "tone": tone}


def compare(prev, cur, quarter=None):
    d = {"previous_file": Path(prev["source_file"]).name, "current_file": Path(cur["source_file"]).name,
         "sheet": cur["sheet"], "quarter": quarter, "themes": [], "changes": [], "unmatched": []}
    ts = {"compared": 0, "moved": 0, "unchanged": 0, "declined": 0, "new": 0, "removed": 0}
    as_ = {"new": 0, "removed": 0, "completed": 0, "progressed": 0, "regressed": 0, "status": 0, "reworded": 0}
    ch = d["changes"]
    years = sorted(set(t["year"] for t in prev["themes"]) | set(t["year"] for t in cur["themes"]))
    for y in years:
        pairs = _match([t for t in prev["themes"] if t["year"] == y], [t for t in cur["themes"] if t["year"] == y], lambda t: t["theme"] or "")
        for pt, ct in pairs:
            name = (ct or pt)["theme"] or "(no theme name)"
            e = {"year": y, "theme": name,
                 "prev_pct": theme_pct(pt) if pt else None, "cur_pct": theme_pct(ct) if ct else None,
                 "prev_status": vd.derive_status([r["status"] for r in pt["rows"]]) if pt else None,
                 "cur_status": vd.derive_status([r["status"] for r in ct["rows"]]) if ct else None,
                 "prev_actions": len(pt["rows"]) if pt else 0, "cur_actions": len(ct["rows"]) if ct else 0, "delta": None}
            if pt and ct:
                e["delta"] = e["cur_pct"] - e["prev_pct"]
                e["state"] = "moved" if e["delta"] > 0 else ("declined" if e["delta"] < 0 else "unchanged")
                ts["compared"] += 1
                ts[e["state"]] += 1
                if pt["theme"] != ct["theme"]:
                    e["renamed_from"] = pt["theme"]
                    ch.append(_row(y, name, f"Key theme renamed from '{pt['theme']}'", "wording", "same"))
                if e["prev_status"] != e["cur_status"]:
                    tone = "up" if e["cur_status"] in ("Completed", "In progress") and e["prev_status"] in ("Not started", "On hold") else ("down" if e["cur_status"] == "Delayed" else "same")
                    ch.append(_row(y, name, "Key theme status", f"{e['prev_status']} → {e['cur_status']}", tone))
                for pa, ca in _match(pt["rows"], ct["rows"], lambda r: r["action"], 0.80):
                    if pa and ca:
                        pp, cp, ps, cs = _p(pa["progress"]), _p(ca["progress"]), pa["status"], ca["status"]
                        if pa["action"] != ca["action"]:
                            as_["reworded"] += 1
                            ch.append(_row(y, name, f"Action reworded: {ca['action']}", f"was: {pa['action']}", "same"))
                        if cs == "Completed" and ps != "Completed":
                            as_["completed"] += 1
                            ch.append(_row(y, name, f"Completed: {ca['action']}", f"{ps or '—'} → Completed", "up"))
                        elif (pp or 0) != (cp or 0):
                            up = (cp or 0) > (pp or 0)
                            as_["progressed" if up else "regressed"] += 1
                            ch.append(_row(y, name, f"Progress: {ca['action']}", f"{pp if pp is not None else '—'}% → {cp if cp is not None else '—'}%", "up" if up else "down"))
                        elif ps != cs:
                            as_["status"] += 1
                            tone = "up" if cs in ("In progress", "Completed") else ("down" if cs == "Delayed" else "same")
                            ch.append(_row(y, name, f"Status: {ca['action']}", f"{ps or '—'} → {cs or '—'}", tone))
                    elif ca:
                        as_["new"] += 1
                        ch.append(_row(y, name, f"New action: {ca['action']}", f"{ca['status'] or '—'} / {_p(ca['progress']) if ca['progress'] is not None else 0}%", "new"))
                    else:
                        as_["removed"] += 1
                        ch.append(_row(y, name, f"Action removed: {pa['action']}", f"was {pa['status'] or '—'} / {_p(pa['progress']) if pa['progress'] is not None else 0}%", "gone"))
            elif ct:
                e["state"] = "new"
                ts["new"] += 1
                d["unmatched"].append(f"new theme {y}: {name} ({len(ct['rows'])} actions)")
                ch.append(_row(y, name, f"New key theme with {len(ct['rows'])} actions", f"{e['cur_pct']}% | {e['cur_status']}", "new"))
            else:
                e["state"] = "removed"
                ts["removed"] += 1
                d["unmatched"].append(f"removed theme {y}: {name} ({len(pt['rows'])} actions in the previous tracker)")
                ch.append(_row(y, name, f"Key theme no longer in the tracker ({len(pt['rows'])} actions)", f"was {e['prev_pct']}% | {e['prev_status']}", "gone"))
            d["themes"].append(e)

    # ---- KPIs
    ks = {"changed_cells": 0, "new": 0, "removed": 0}
    q = quarter or cur.get("quarter") or "Q2"
    critical_deltas = []
    _pa = prev["kpis"]["critical"] + prev["kpis"]["enabling"]
    _ca = cur["kpis"]["critical"] + cur["kpis"]["enabling"]
    _kind = {id(k): ("critical" if k in cur["kpis"]["critical"] else "enabling") for k in _ca}
    if True:
        for pk, ck in _match(_pa, _ca, lambda k: k["metric"]):
            kind = _kind.get(id(ck), "previous") if ck is not None else "previous"
            if pk and ck:
                if kind == "critical":
                    pv = pk["quarters"].get(q, {}).get("actual")
                    cv = ck["quarters"].get(q, {}).get("actual")
                    tgt = ck["quarters"].get(q, {}).get("target")
                    pn, cn, tn = vd.parse_number(pv), vd.parse_number(cv), vd.parse_number(tgt)
                    delta_val = (cn - pn) if (cn is not None and pn is not None) else None
                    lower_better = ck.get("lower_is_better", False)
                    improved = None
                    if delta_val is not None:
                        improved = (delta_val < 0) if lower_better else (delta_val > 0)
                    target_achieved = None
                    if cn is not None and tn is not None:
                        target_achieved = (cn <= tn) if lower_better else (cn >= tn)
                    critical_deltas.append({
                        "metric": ck["metric"],
                        "kind_of_value": ck["kind_of_value"],
                        "target": tgt,
                        "target_num": tn,
                        "prev_actual": pv,
                        "prev_num": pn,
                        "cur_actual": cv,
                        "cur_num": cn,
                        "delta": delta_val,
                        "improved": improved,
                        "target_achieved": target_achieved,
                        "lower_is_better": lower_better
                    })
                fields = [("baseline", pk["baseline"], ck["baseline"], "2025 baseline")]
                for yy in ck["targets"]:
                    fields.append((f"target_{yy}", pk["targets"].get(yy), ck["targets"].get(yy), f"{yy} target"))
                for qq in ck["quarters"]:
                    fields.append((f"{qq}_target", pk["quarters"].get(qq, {}).get("target"), ck["quarters"][qq]["target"], f"{qq} target"))
                    fields.append((f"{qq}_actual", pk["quarters"].get(qq, {}).get("actual"), ck["quarters"][qq]["actual"], f"{qq} actual"))
                if (pk["committed"] or "") != (ck["committed"] or ""):
                    fields.append(("committed", pk["committed"], ck["committed"], "Committed action"))
                for key, pv, cv, label in fields:
                    if pv == cv:
                        continue
                    ks["changed_cells"] += 1
                    tone = "same"
                    if key.endswith("_actual"):
                        pn, cn = vd.parse_number(pv), vd.parse_number(cv)
                        if pn is not None and cn is not None and pn != cn:
                            better = cn < pn if ck.get("lower_is_better") else cn > pn
                            tone = "up" if better else "down"
                        elif pv is None and cv is not None:
                            tone = "new"
                    pv_s = pv if isinstance(pv, str) or pv is None else vd.fmt_value(pv, ck["kind_of_value"])
                    cv_s = cv if isinstance(cv, str) or cv is None else vd.fmt_value(cv, ck["kind_of_value"])
                    ch.append(_row("KPI", ck["metric"], f"{label} ({kind})", f"{pv_s if pv_s is not None else '—'} → {cv_s if cv_s is not None else '—'}", tone))
            elif ck:
                ks["new"] += 1
                ch.append(_row("KPI", ck["metric"], f"New {kind} KPI", "added", "new"))
            else:
                ks["removed"] += 1
                ch.append(_row("KPI", pk["metric"], f"{kind.title()} KPI no longer in the tracker", "removed", "gone"))
    d["critical_kpi_deltas"] = critical_deltas

    # ---- dependencies (asks placed on this function by other functions)
    ds = {"previous": len(prev["dependencies"]), "current": len(cur["dependencies"]), "new": 0, "removed": 0, "completed": 0, "moved": 0}
    for pd_, cd_ in _match(prev["dependencies"], cur["dependencies"], lambda x: f"{x['function']}|{x['action'] or ''}"):
        if pd_ and cd_:
            pp, cp = _p(pd_["progress"]), _p(cd_["progress"])
            if cd_["status"] == "Completed" and pd_["status"] != "Completed":
                ds["completed"] += 1
                ch.append(_row("Dep.", f"{cd_['function']}: {cd_['theme'] or ''}", f"Dependency completed: {cd_['action']}", f"{pd_['status'] or '—'} → Completed", "up"))
            elif (pp or 0) != (cp or 0) or pd_["status"] != cd_["status"]:
                ds["moved"] += 1
                up = (cp or 0) >= (pp or 0) and cd_["status"] != "Delayed"
                ch.append(_row("Dep.", f"{cd_['function']}: {cd_['theme'] or ''}", f"Dependency: {cd_['action']}", f"{pp if pp is not None else 0}% {pd_['status'] or '—'} → {cp if cp is not None else 0}% {cd_['status'] or '—'}", "up" if up else "down"))
        elif cd_:
            ds["new"] += 1
            ch.append(_row("Dep.", f"{cd_['function']}: {cd_['theme'] or ''}", f"New dependency: {cd_['action']}", f"{cd_['status'] or '—'}", "new"))
        else:
            ds["removed"] += 1
            ch.append(_row("Dep.", f"{pd_['function']}: {pd_['theme'] or ''}", f"Dependency removed: {pd_['action']}", f"was {pd_['status'] or '—'}", "gone"))

    # ---- structured per-KPI delta (reporting-quarter actual), used by the delta chips and columns
    d["kpis"] = []
    q = quarter
    _prev_all = prev["kpis"]["critical"] + prev["kpis"]["enabling"]
    _cur_all = cur["kpis"]["critical"] + cur["kpis"]["enabling"]
    _kind_of = {id(k): ("critical" if k in cur["kpis"]["critical"] else "enabling") for k in _cur_all}
    if True:
        for pk, ck in _match(_prev_all, _cur_all, lambda k: k["metric"]):
            if ck is None:
                continue
            kind = _kind_of.get(id(ck), "enabling")
            pa = pk["quarters"].get(q, {}).get("actual") if (pk and q) else None
            ca = ck["quarters"].get(q, {}).get("actual") if q else None
            pn, cn = vd.parse_number(pa), vd.parse_number(ca)
            e = {"metric": ck["metric"], "kind": kind, "in_previous": pk is not None,
                 "prev": pa, "cur": ca, "delta": None, "direction": "none", "text": "No previous value"}
            if pk is None:
                e["text"] = "New this cycle"
                e["direction"] = "new"
            elif pn is None or cn is None:
                e["text"] = "No comparable value"
            elif pn == cn:
                e["text"] = "No change"
                e["direction"] = "flat"
                e["delta"] = 0
            else:
                diff = cn - pn
                e["delta"] = diff
                better = diff < 0 if ck.get("lower_is_better") else diff > 0
                e["direction"] = "up" if better else "down"
                kindv = ck["kind_of_value"]
                if kindv == "percent":
                    scale = 100 if (abs(pn) <= 1 and abs(cn) <= 1) else 1
                    e["text"] = f"{diff * scale:+.1f} pts"
                elif kindv == "currency":
                    e["text"] = f"{'+' if diff > 0 else ''}€{abs(diff):.2f}M" if diff < 0 else f"+€{diff:.2f}M"
                    if diff < 0:
                        e["text"] = f"-€{abs(diff):.2f}M"
                else:
                    e["text"] = f"{diff:+g}"
            d["kpis"].append(e)

    # ---- dependency status counts and per-function progress, previous vs current
    def dep_counts(m):
        o = {}
        for x in m["dependencies"]:
            o[x["status"] or "Not recorded"] = o.get(x["status"] or "Not recorded", 0) + 1
        return o

    def dep_by_function(m):
        o = {}
        for x in m["dependencies"]:
            o.setdefault(x["function"], []).append(x["progress"] if x["progress"] is not None else 0.0)
        return {k: {"avg": round(100 * sum(v) / len(v)), "n": len(v)} for k, v in o.items()}
    d["dep_status"] = {"prev": dep_counts(prev), "cur": dep_counts(cur),
                       "prev_total": len(prev["dependencies"]), "cur_total": len(cur["dependencies"])}
    pf, cf = dep_by_function(prev), dep_by_function(cur)
    d["dep_functions"] = []
    for f in sorted(set(pf) | set(cf), key=str.lower):
        p_, c_ = pf.get(f), cf.get(f)
        d["dep_functions"].append({"function": f, "prev": p_["avg"] if p_ else None, "cur": c_["avg"] if c_ else None,
                                   "n_prev": p_["n"] if p_ else 0, "n_cur": c_["n"] if c_ else 0,
                                   "delta": (c_["avg"] - p_["avg"]) if (p_ and c_) else None})

    order = {"up": 0, "down": 1, "new": 2, "gone": 3, "same": 4}
    ch.sort(key=lambda c: (0 if c["year"].isdigit() else (1 if c["year"] == "KPI" else 2), c["year"], c["theme"].lower(), order.get(c["tone"], 9)))
    d["theme_summary"], d["action_summary"], d["kpi_summary"], d["dependency_summary"] = ts, as_, ks, ds
    return d


# ----------------------------------------------------------------------------- slides
def _short(s, n=LABEL_MAX):
    s = s or ""
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def add_delta_slides(prs, view, s2, s4, extra_added, note):
    d = view["delta"]
    ts = d["theme_summary"]
    idx4 = prs.slides.index(s4)

    # --- slide A: cards + native bar chart (duplicate of the commitments slide, table replaced by the chart)
    sa = vd.duplicate_slide(prs, s2, idx4)
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 15"), "MONTH-ON-MONTH")
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 1"), "Key theme movement vs previous tracker")
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 5"), str(ts["compared"]))
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 6"), "Key themes compared")
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 8"), str(ts["declined"] + ts["unchanged"] + ts["removed"]))
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 9"), "Themes without forward movement")
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 11"), str(ts["moved"] + ts["new"]))
    vd.set_shape_text(vd.shape_by_name(sa, "Rectangle 12"), "Themes moved forward or newly added")
    tbl = vd.shape_by_name(sa, "Table 28")
    tbl._element.getparent().remove(tbl._element)

    themes = d["themes"]
    cd = CategoryChartData()
    cd.categories = [f"{t['year']} · {_short(t['theme'])}" for t in themes] or ["No key themes"]
    cd.add_series(f"Previous ({d['previous_file']})", [t["prev_pct"] if t["prev_pct"] is not None else 0 for t in themes] or [0])
    cd.add_series(f"Current ({d['current_file']})", [t["cur_pct"] if t["cur_pct"] is not None else 0 for t in themes] or [0])
    gf = sa.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.44), Inches(3.35), Inches(12.46), Inches(3.75), cd)
    gf.name = "Chart Delta"
    chart = gf.chart
    chart.has_title = False
    chart.font.name = "Helvetica Neue"
    chart.font.size = Pt(9)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(9)
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.overlap = -12
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0"%"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(8)
    prev_s, cur_s = plot.series[0], plot.series[1]
    prev_s.format.fill.solid()
    prev_s.format.fill.fore_color.rgb = PREV_COLOUR
    cur_s.format.fill.solid()
    cur_s.format.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    cat = chart.category_axis
    cat.reverse_order = True
    cat.tick_labels.font.size = Pt(9)
    cat.has_major_gridlines = False
    cat.format.line.color.rgb = GRID_COLOUR
    val = chart.value_axis
    val.minimum_scale, val.maximum_scale, val.major_unit = 0, 100, 25
    val.has_major_gridlines = True
    val.major_gridlines.format.line.color.rgb = GRID_COLOUR
    val.tick_labels.number_format = '0"%"'
    val.tick_labels.number_format_is_linked = False
    val.tick_labels.font.size = Pt(8)
    val.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    val.format.line.fill.background()
    vd.set_notes(sa, note)

    # --- slide B: "what changed" table (themes-slide prototype: 4 columns)
    sb = vd.duplicate_slide(prs, s4, idx4 + 1)
    vd.set_shape_text(vd.shape_by_name(sb, "Rectangle 6"), "MONTH-ON-MONTH")
    title = "What changed since the previous tracker"
    vd.set_shape_text(vd.shape_by_name(sb, "Rectangle 1"), title)
    rows = [[c["year"], c["theme"], c["change"], c["movement"]] for c in d["changes"]]
    statuses = [TONE_PALETTE.get(c["tone"]) for c in d["changes"]]
    if not rows:
        rows = [["", "", f"No changes between {d['previous_file']} and {d['current_file']}", ""]]
        statuses = [None]
        extra_added.append("placeholder text 'No changes between the two trackers' on the delta table (no changes found)")
    header = ["Year", "Key theme / KPI", "Change", "Previous → Current"]
    geo = vd._table_geometry(vd.shape_by_name(sb, "Table 10"))
    groups = vd.pack_rows(rows, geo)
    vd.fill_section(prs, sb, "Table 10", rows, statuses, 3, header, groups, title, extra_added, note=note)
    extra_added.append("month-on-month delta slides added on request: 'Key theme movement vs previous tracker' (native chart) and 'What changed since the previous tracker'")


# ----------------------------------------------------------------------------- summary
def summary_lines(d):
    ts, as_, ks, ds = d["theme_summary"], d["action_summary"], d["kpi_summary"], d["dependency_summary"]
    L = ["## Month-on-month delta",
         f"- Previous tracker: `{d['previous_file']}`  ->  current tracker: `{d['current_file']}` (sheet '{d['sheet']}')",
         "- Matching: items paired by normalised name (exact first, then fuzzy >= 0.85 for themes/KPIs/dependencies, >= 0.80 for actions); unmatched items are listed as new/removed.",
         f"- Key themes: {ts['compared']} compared, {ts['moved']} moved forward, {ts['unchanged']} unchanged, {ts['declined']} declined, {ts['new']} new, {ts['removed']} removed.",
         f"- Actions: {as_['completed']} completed, {as_['progressed']} progressed, {as_['regressed']} regressed, {as_['status']} status changes, {as_['new']} new, {as_['removed']} removed, {as_['reworded']} reworded.",
         f"- KPI cells changed: {ks['changed_cells']} ({ks['new']} new KPIs, {ks['removed']} removed).",
         f"- Dependencies: {ds['previous']} -> {ds['current']} rows; {ds['completed']} completed, {ds['moved']} moved, {ds['new']} new, {ds['removed']} removed.",
         f"- Change rows on the 'What changed' slide: {len(d['changes'])}. Theme progress is derived exactly as on the themes slide (Key Theme Progress column, else average of action progress).",
         f"- Chart category labels are shortened to {LABEL_MAX} characters; full theme names are on the themes and change slides."]
    for u in d["unmatched"]:
        L.append(f"  - {u}")
    return L
