# -*- coding: utf-8 -*-
"""Build the standard VCP review deck for every function sheet of a tracker.

    python vcp_agent/build_all.py --xlsx <current> [--previous <older>] [--governance <gov>] [--outdir output/all]

Every function gets the same slide set; functions whose sheet holds no themes, KPIs or dependencies
still get the slides, with the generator's own "not recorded in the sheet" placeholders.
"""
import argparse
import sys
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import vcp_deck as vd  # noqa: E402

SKIP = {"KEY RISKS", "Portfolio", "Marketing_Aug26"}   # no VCP tracker layout / superseded by Marketing_New
DECISIONS = [
    "Update PPT template to include previous meeting action tracker, KPI risk status, owner, timeline, support required, and root-cause tracking.",
    "Incorporate linkage between themes/actions and KPI impact within the review format.",
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--previous")
    ap.add_argument("--governance")
    ap.add_argument("--outdir", default="output/all")
    ap.add_argument("--quarter", default="Q2")
    ap.add_argument("--only", nargs="*")
    ap.add_argument("--render", action="store_true")
    ap.add_argument("--with-change-table", action="store_true")
    ap.add_argument("--with-governance-slides", action="store_true")
    a = ap.parse_args()
    tr = vd.Tracker(a.xlsx)
    funcs = [n for n, s in tr.sheets() if s == "visible" and n not in SKIP]
    if a.only:
        funcs = [f for f in funcs if f in a.only]
    rows = []
    for f in funcs:
        try:
            res = vd.generate(a.xlsx, f, outdir=a.outdir, quarter=a.quarter, allow_extra=True,
                              impact_mode="keep", previous=a.previous, governance=a.governance,
                              decisions=DECISIONS, with_change_table=a.with_change_table,
                              with_governance_slides=a.with_governance_slides)
            if res["questions"]:
                rows.append((f, "QUESTIONS", 0, False, "; ".join(q[:110] for q in res["questions"])))
                continue
            from pptx import Presentation
            n = len(Presentation(str(res["pptx"])).slides)
            rows.append((f, "built", n, res["reconciled"], Path(res["pptx"]).name))
            if a.render:
                vd.render(res["pptx"], Path(a.outdir) / "render")
        except Exception as e:  # noqa: BLE001
            rows.append((f, "ERROR", 0, False, f"{type(e).__name__}: {e}"))
            traceback.print_exc(limit=3)
    print("\n=== build_all summary ===")
    for f, st, n, ok, note in rows:
        print(f"  {f:22s} {st:9s} slides={n:3d} reconciled={str(ok):5s} {note[:150]}")
    bad = [r for r in rows if r[1] != "built" or not r[3]]
    print(f"\n{len(rows) - len(bad)}/{len(rows)} functions built and reconciled")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
