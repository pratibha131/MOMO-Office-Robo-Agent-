# -*- coding: utf-8 -*-
"""
Renderer for the 'Monthly VCP report' template (vcp_agent/template/CS_VCP_Monthly_Report.pptx).

The template is the single source of design. Every text element is either a verbatim Excel value,
a template label, or a documented derivation (counts, ratios, gaps) - the same rules as the
original renderer. Narrative commentary that the template carried as an example (remarks,
management actions, engine names) is replaced by data-derived statements; nothing is invented.

Slide map (template slide -> content)
  1 cover                       title
  2 Critical VCP commitments    cards, target-normalised chart (per-KPI colours + value labels), remarks per KPI
  3 Enabling KPIs               cards, table
  4 Themes and actions          table
  5 Interdependencies           cards, leadership focus, dependency register (one shape row per action), source line
  6 VCP impact linkage          cards, driver map (one row per critical KPI), outcome trajectory chart, iGM chip
  7 RCA and sheet performance   RCA blocks per red KPI, sheet performance dashboard
"""
import copy

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import vcp_deck as vd

# ----------------------------------------------------------------------------- template palette (from the XML)
C = {"navy": "103B66", "blue": "1769C2", "ink": "152033", "muted": "5B6678", "line": "D7DEE8", "green": "238A4B",
     "red": "C9363E", "amber": "A36B00", "teal": "188B84", "blue_bg": "E6F0FB", "green_bg": "DDF1E4", "grey_bg": "EDF0F4",
     "red_bg": "F9DFE1", "amber_bg": "FFF1CE", "row_bg": "F4F7FA", "bar_bg": "E8EDF3"}
# status pills / status table cells (fill, font)
STATUS2 = {"In progress": {"fill": "E6F0FB", "font": "1769C2"}, "Completed": {"fill": "DDF1E4", "font": "238A4B"},
           "Not started": {"fill": "EDF0F4", "font": "5B6678"}, "Delayed": {"fill": "F9DFE1", "font": "C9363E"},
           "On hold": {"fill": "EDF0F4", "font": "5B6678"}, "No data": {"fill": "EDF0F4", "font": "5B6678"}}
TONE2 = {"up": "Completed", "down": "Delayed", "new": "In progress", "gone": "No data", "same": "No data"}
SERIES_COLOURS = ["1769C2", "188B84", "A36B00", "C9363E", "238A4B"]

THEME_DELTA_WIDTHS = [0.62, 2.62, 5.56, 1.50, 1.65]   # themes table when the delta column is added
REGISTER_ROWS = 16          # register rows per slide (template geometry)
REGISTER_PITCH = 0.1967     # inches between register rows
RCA_MAX = 6                 # RCA blocks per slide (4 at template pitch, up to 6 compressed)
REMARKS_MAX = 5             # remark blocks per slide
DRIVER_MAX = 4              # driver-map rows per slide


TONES = {"up": {"fill": "DDF1E4", "font": "238A4B", "mark": "▲"},
         "down": {"fill": "F9DFE1", "font": "C9363E", "mark": "▼"},
         "flat": {"fill": "EDF0F4", "font": "5B6678", "mark": "–"},
         "new": {"fill": "E6F0FB", "font": "1769C2", "mark": "✦"},
         "none": {"fill": "EDF0F4", "font": "5B6678", "mark": "–"}}


def chip(slide, x, y, w, h, text, tone, size=8.0):
    """Rounded status chip in the template's own style (used for the month-on-month deltas)."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    t = TONES.get(tone, TONES["none"])
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = f"Delta chip {text[:24]}"
    try:
        sh.adjustments[0] = 0.28
    except (IndexError, ValueError):
        pass
    fill(sh, t["fill"])
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(int(Inches(0.05)))
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"{t['mark']}  {text}"
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.name = "Arial"
    run.font.color.rgb = rgb(t["font"])
    return sh


def chip_row(slide, items, y, x0=0.65, x1=7.95, h=0.42, size=8.0, gap=0.16):
    """Lay out n chips evenly between x0 and x1."""
    n = max(len(items), 1)
    w = min(2.40, (x1 - x0 - (n - 1) * gap) / n)
    for i, (text, tone) in enumerate(items):
        chip(slide, x0 + i * (w + gap), y, w, h, text, tone, size)


def delta_mark(e):
    """(text, tone) for one KPI delta entry: value carries an increase/dip arrow and a red/green tone."""
    if e is None:
        return "No previous value", "none"
    d = e.get("direction")
    if d == "new":
        return "New this cycle", "new"
    if d == "flat":
        return "0 · no change", "flat"
    if d in ("up", "down"):
        return e["text"], d
    return e.get("text") or "No comparable value", "none"


def theme_delta_mark(e):
    """(text, tone) for one key-theme progress delta."""
    if e is None or e.get("prev_pct") is None or e.get("delta") is None:
        return "No previous value", "none"
    d = e["delta"]
    if d == 0:
        return f"0 pts · no change (was {e['prev_pct']}%)", "flat"
    return f"{d:+d} pts (was {e['prev_pct']}%)", ("up" if d > 0 else "down")


def marked(text, tone):
    return f"{TONES.get(tone, TONES['none'])['mark']}  {text}"


def tone_palette(tone):
    t = TONES.get(tone, TONES["none"])
    return {"fill": t["fill"], "font": t["font"]}


def kpi_delta_map(view):
    return {k["metric"]: k for k in (view.get("delta") or {}).get("kpis", [])}


def theme_delta_map(view):
    return {(str(t["year"]), t["theme"]): t for t in (view.get("delta") or {}).get("themes", [])}


def is_v2(prs):
    try:
        return any(s.name == "Chart" for s in list(prs.slides)[1].shapes)
    except IndexError:
        return False


# ----------------------------------------------------------------------------- shape helpers
def sh(slide, name):
    return vd.shape_by_name(slide, name)


def txt(slide, name, text):
    vd.set_shape_text(sh(slide, name), text)


def rgb(hexv):
    return RGBColor.from_string(hexv)


def fill(shape, hexv):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(hexv)


def color(shape, hexv):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = rgb(hexv)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def _next_id(slide):
    return max([s.shape_id for s in slide.shapes] + [1]) + 1


def clone(slide, shape, dy=0.0, dx=0.0, name=None):
    el = copy.deepcopy(shape._element)
    cnv = el.find(".//" + qn("p:cNvPr"))
    cnv.set("id", str(_next_id(slide)))
    if name:
        cnv.set("name", name)
    slide.shapes._spTree.append(el)
    new = slide.shapes[-1]
    if dx:
        new.left = Emu(int(new.left + Inches(dx)))
    if dy:
        new.top = Emu(int(new.top + Inches(dy)))
    return new


def shapes_in_band(slide, y0, y1, exclude=()):
    """Shapes whose top edge lies in [y0, y1] inches (used to clear template example rows)."""
    out = []
    for s in list(slide.shapes):
        if s.name in exclude:
            continue
        top = s.top / 914400
        if y0 <= top <= y1:
            out.append(s)
    return out


def set_page_numbers(prs):
    content = list(prs.slides)[1:]
    for i, s in enumerate(content, 1):
        for shp in s.shapes:
            if shp.name == "Rectangle 4" and shp.has_text_frame and "report |" in shp.text_frame.text:
                vd.set_shape_text(shp, f"Monthly VCP report | {i}/{len(content)}")


# ----------------------------------------------------------------------------- chart helper
def update_chart(gf, categories, series, series_colours=None, point_colours=None, labels=None, label_colours=None,
                 label_pt=7.13, y_max=None, y_min=None):
    """Rewrite a template chart's data while keeping its styling. labels[si][pi] = custom text per point."""
    chart = gf.chart
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series:
        cd.add_series(name, vals)
    for ser in chart.plots[0].series:                      # drop per-point overrides of the example data
        el = ser._element
        for tag in ("c:dPt", "c:dLbls"):
            for x in el.findall(qn(tag)):
                el.remove(x)
    chart.replace_data(cd)
    plot = chart.plots[0]
    for si, ser in enumerate(plot.series):
        if series_colours and si < len(series_colours) and series_colours[si]:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(series_colours[si])
        for pi in range(len(categories)):
            pt = ser.points[pi]
            pc = point_colours[si][pi] if point_colours and point_colours[si] else None
            if pc:
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(pc)
            if labels and labels[si] and labels[si][pi] is not None:
                dl = pt.data_label
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
                tf = dl.text_frame
                tf.text = str(labels[si][pi])
                run = tf.paragraphs[0].runs[0]
                run.font.size = Pt(label_pt)
                run.font.bold = True
                lc = (label_colours[si][pi] if label_colours and label_colours[si] else None) or pc or (series_colours[si] if series_colours and si < len(series_colours) else None)
                if lc:
                    run.font.color.rgb = rgb(lc)
    if y_max is not None:
        chart.value_axis.maximum_scale = y_max
    if y_min is not None:
        chart.value_axis.minimum_scale = y_min


# ----------------------------------------------------------------------------- derived texts
def kpi_state(k, q, issues, sheet):
    """(signal, ratio actual/target or None, formatted target, formatted actual, gap text)"""
    sig, _ = vd.signal_for(k, q, issues, sheet)
    d = k["quarters"].get(q, {"target": None, "actual": None})
    t, a = vd.parse_number(d["target"]), vd.parse_number(d["actual"])
    ratio = (a / t) if (t not in (None, 0) and a is not None) else None
    ft = vd.fmt_value(d["target"], k["kind_of_value"]) if d["target"] is not None else None
    fa = vd.fmt_value(d["actual"], k["kind_of_value"]) if d["actual"] is not None else None
    gap = vd.delta_text(k, q).strip("()") if (t is not None and a is not None) else None
    return sig, ratio, ft, fa, gap


def gap_abs(k, gap):
    """'-0.30M' -> '€0.30M', '-1.6 pts' -> '1.6 pts' (magnitude for prose; pills keep the signed form)."""
    g = (gap or "").lstrip("+-").strip()
    if k["kind_of_value"] == "currency" and g and not g.startswith("€"):
        g = "€" + g
    return g


def remark_for(k, q, sig, ft, fa, gap):
    committed = k["committed"] or "not captured in the sheet"
    if sig == "Below target":
        head = f"{gap_abs(k, gap)} below the {q} target ({fa} vs {ft})."
    elif sig == "On / above target":
        head = (f"At the {q} target ({fa})." if fa == ft else f"{gap.lstrip('+') if gap else ''} above the {q} target ({fa} vs {ft}).".replace("  ", " "))
    else:
        head = f"No {q} target/actual pair in the sheet."
    return f"{head} Committed action: {committed}"


def short_gap(gap):
    return gap or "—"


# ----------------------------------------------------------------------------- planning (capacity)
def plan_v2(view, template, impact_mode=None):
    """Groups per section (same shape as plan_tables) + overflow strings + questions."""
    prs = vd.Presentation(str(template))
    slides = list(prs.slides)
    plan, over, qs = {}, [], []
    geo_e = vd._table_geometry(sh(slides[2], "Table 54"))
    geo_t = vd._table_geometry(sh(slides[3], "Table 13"))
    if view.get("delta"):        # the themes table gains a column, so pack against the narrower widths
        geo_t = dict(geo_t, widths=THEME_DELTA_WIDTHS)   # keep the template point size: pack_rows floors at MIN_BODY_PT
    for key, geo, rows in (("enabling", geo_e, view["enabling"]["rows"]), ("themes", geo_t, [t["cells"] for t in view["themes"]])):
        chosen, groups = geo["pt"], [[]]
        if rows:
            for pt in (geo["pt"], geo["pt"] - 1, geo["pt"] - 2):
                if pt < vd.MIN_BODY_PT:
                    break
                groups, chosen = vd.pack_rows(rows, geo, pt=pt), pt
                if len(groups) == 1:
                    break
            if len(groups) > 1:
                groups, chosen = vd.pack_rows(rows, geo), geo["pt"]
        plan[key] = {"groups": groups, "pt": chosen, "orig_pt": geo["pt"]}
        if len(groups) > 1:
            over.append(f"{key}: {len(rows)} rows need {len(groups)} slides")

    def chunk(n, per):
        return [list(range(i, min(i + per, n))) for i in range(0, n, per)] or [[]]
    n_crit = len(view["critical"]["rows"])
    n_dep = len(view["dep_register"])
    n_red = len(view["rca"]["rows"])
    for key, n, per in (("critical", n_crit, REMARKS_MAX), ("dependencies", n_dep, REGISTER_ROWS), ("rca", n_red, RCA_MAX), ("impact", n_crit, DRIVER_MAX)):
        if key == "impact" and impact_mode == "drop" and not view.get("outcomes"):
            plan[key] = {"groups": [[]], "pt": None, "orig_pt": None}
            continue
        g = chunk(n, per)
        plan[key] = {"groups": g, "pt": None, "orig_pt": None}
        if len(g) > 1:
            over.append(f"{key}: {n} rows need {len(g)} slides")
    if not view.get("outcomes") and impact_mode is None:
        qs.append("The sheet has no VCP outcome table (a 'Metric | Year | Q1..Q4 | FY' block), so the 'VCP impact linkage' slide has no "
                  "trajectory data. Choose keep (driver map from the critical KPIs, trajectory marked 'Not quantified') or drop (remove the slide).")
    return plan, over, qs


# ----------------------------------------------------------------------------- slide builders
def build_critical(slide, view, model, extra_added):
    q = view["quarter"]
    crit = model["kpis"]["critical"]
    sheet = model["sheet"]
    states = [kpi_state(k, q, model["issues"], sheet) for k in crit]
    charted = [(k, st) for k, st in zip(crit, states) if st[1] is not None]
    below = [k["metric"] for k, st in zip(crit, states) if st[0] == "Below target"]
    not_charted = [k["metric"] for k, st in zip(crit, states) if st[1] is None]
    txt(slide, "Rectangle 2", f"{q} target vs actual for critical KPIs")
    sub = "Target bars are blue. Actual bars turn green when achieved and red when below target."
    if not_charted:
        sub += f" Not charted (no {q} target/actual pair): {', '.join(not_charted)}."
    txt(slide, "Rectangle 3", sub)
    txt(slide, "Rectangle 12", f"{len(crit):02d}")
    txt(slide, "Rectangle 13", "Total Critical KPIs")
    txt(slide, "Rectangle 52", f"{len(below)}/{len(crit)}")
    txt(slide, "Rectangle 54", f"Critical KPIs below {q} target")
    txt(slide, "Rectangle 56", (", ".join(below) + " below target") if below else f"No critical KPI below the {q} target")
    if not below:
        fill(sh(slide, "Rectangle 50"), C["green"])
        color(sh(slide, "Rectangle 52"), C["green"])
    txt(slide, "Rectangle 16", f"{q} target and actual by critical KPI")
    txt(slide, "Rectangle 18", f"{q} target")
    # chart: target normalised to 100 %, actual as a share of target; colours by signal
    gf = sh(slide, "Chart")
    if charted:
        cats = [k["metric"] for k, _ in charted]
        ratios = [st[1] for _, st in charted]
        cols = [C["green"] if st[0] == "On / above target" else C["red"] for _, st in charted]
        update_chart(gf, cats, [(f"{q} target", [1.0] * len(cats)), ("Actual", ratios)],
                     series_colours=[C["blue"], C["red"]], point_colours=[None, cols],
                     labels=[[st[2] for _, st in charted], [st[3] for _, st in charted]],
                     label_colours=[[C["blue"]] * len(cats), cols], y_max=max(1.12, max(ratios) * 1.15))
        va = gf.chart.value_axis
        va.tick_labels.number_format = "0%"
        va.tick_labels.number_format_is_linked = False
    else:
        update_chart(gf, ["No KPI with target and actual"], [(f"{q} target", [0]), ("Actual", [0])], series_colours=[C["blue"], C["red"]])
        extra_added.append(f"placeholder chart category 'No KPI with target and actual' on the critical KPI slide (no {q} target/actual pairs)")
    # remarks: one block per critical KPI (template has 3 example blocks)
    proto = {n: sh(slide, n) for n in ("Rectangle 27", "Rectangle 28", "Rectangle 29")}
    for n in ("Rectangle 30", "Rectangle 31", "Rectangle 32", "Rectangle 33", "Rectangle 34"):
        remove(sh(slide, n))
    n = max(len(crit), 1)
    top0 = 2.27
    pitch = min(1.47, (6.69 - top0) / n)
    for i, (k, st) in enumerate(zip(crit, states)):
        lab = proto["Rectangle 27"] if i == 0 else clone(slide, proto["Rectangle 27"], dy=i * pitch)
        body = proto["Rectangle 28"] if i == 0 else clone(slide, proto["Rectangle 28"], dy=i * pitch)
        vd.set_shape_text(lab, k["metric"])
        col = C["red"] if st[0] == "Below target" else (C["green"] if st[0] == "On / above target" else C["muted"])
        color(lab, col)
        vd.set_shape_text(body, remark_for(k, q, st[0], st[2], st[3], st[4]))
        body.height = Emu(int(Inches(pitch - 0.16)))
        if i < len(crit) - 1:
            sep = proto["Rectangle 29"] if i == 0 else clone(slide, proto["Rectangle 29"], dy=i * pitch)
            sep.top = Emu(int(Inches(top0 + (i + 1) * pitch - 0.08)))
        elif i == 0:
            remove(proto["Rectangle 29"])
    if not crit:
        vd.set_shape_text(proto["Rectangle 27"], "KPIs")
        vd.set_shape_text(proto["Rectangle 28"], "No critical KPI rows in the sheet.")
        remove(proto["Rectangle 29"])
    # month-on-month delta chips under the chart (one per critical KPI)
    if view.get("delta") and crit:
        dm = kpi_delta_map(view)
        panel = sh(slide, "Rectangle: Rounded Corners 15")
        panel.height = Emu(int(Inches(3.30)))
        gf.height = Emu(int(Inches(2.40)))
        items = []
        for k in crit:
            text, tone = delta_mark(dm.get(k["metric"]))
            items.append((f"{k['metric']}: {text}", tone))
        chip_row(slide, items, y=6.36, x0=0.65, x1=7.95, h=0.42, size=7.5)
        extra_added.append(f"month-on-month delta chips added under the critical KPI chart ({len(items)} KPIs)")
    return [st[0] for st in states]


def build_enabling(slide, view, model, plan, extra_added, group_index=0):
    q = view["quarter"]
    enab = model["kpis"]["enabling"]
    ready = sum(1 for k in enab if vd.kpi_latest_quarter(k) == q)
    red = sum(1 for s in view["enabling"]["signals"] if s == "Below target")
    nodata = sum(1 for k in enab if not vd.kpi_has_any_quarter_data(k))
    n = len(enab)
    txt(slide, "Rectangle 2", f"{ready} of {n} enabling KPIs are {q}-ready; {red} below target")
    no_pair = [k["metric"] for k in enab if vd.kpi_latest_quarter(k) != q]
    txt(slide, "Rectangle 3", f"{q}-ready means a {q} target and actual are both present. " + (f"No {q} pair for: {', '.join(no_pair)}." if no_pair else "All enabling KPIs have a target/actual pair."))
    txt(slide, "Rectangle 7", f"{ready}/{n}")
    txt(slide, "Rectangle 8", f"{q}-ready KPIs")
    txt(slide, "Rectangle 12", str(red))
    txt(slide, "Rectangle 13", f"{q} KPIs red")
    txt(slide, "Rectangle 14", f"{red} of {ready} populated KPIs miss target" if ready else "No populated KPI to assess")
    txt(slide, "Rectangle 17", str(nodata))
    txt(slide, "Rectangle 18", "No quarterly data")
    rows, sigs = view["enabling"]["rows"], view["enabling"]["signals"]
    ly = view["last_year"]
    header = ["KPI", "2025 baseline", f"{ly} target", f"{q} target", f"{q} actual", f"{q} signal", "Committed action"]
    idx = plan["enabling"]["groups"][group_index] if rows else []
    tbl = sh(slide, "Table 54")
    delta_col = None
    if view.get("delta") and rows:
        # extra "Δ vs previous" column, inserted before the signal column
        dm = kpi_delta_map(view)
        vd.set_table_columns(tbl, [1.90, 1.15, 1.10, 1.10, 1.10, 1.30, 1.31, 3.50])
        header = ["KPI", "2025 baseline", f"{ly} target", f"{q} target", f"{q} actual",
                  f"Δ {q} actual vs previous", f"{q} signal", "Committed action"]
        delta_col, tones = 5, []
        newrows = []
        for i in idx:
            r = list(rows[i])
            text, tone = delta_mark(dm.get(r[0]))
            newrows.append(r[:5] + [marked(text, tone)] + r[5:])
            tones.append(tone_palette(tone))
        vd.fill_table(tbl, newrows, status_col=6, status_values=[sigs[i] for i in idx], header=header)
        vd.colour_column(tbl, delta_col, tones)
        vd.set_body_font(tbl, 10)
        extra_added.append("enabling KPI table: added the 'delta vs previous' column")
    elif rows:
        vd.fill_table(tbl, [rows[i] for i in idx], status_col=5, status_values=[sigs[i] for i in idx], header=header)
        p = plan["enabling"]
        if p["pt"] and p["orig_pt"] and p["pt"] != p["orig_pt"] and delta_col is None:
            vd.set_body_font(tbl, p["pt"])
            if group_index == 0:
                extra_added.append(f"body font reduced from {p['orig_pt']:g}pt to {p['pt']:g}pt on the enabling table to keep it on one slide")
    if not rows:
        vd.fill_table(tbl, [["No enabling KPI rows recorded in the sheet", "", "", "", "", "", ""]], header=header)
        extra_added.append("placeholder text 'No enabling KPI rows recorded in the sheet' on the enabling table")


def build_themes(slide, view, plan, extra_added, group_index=0):
    y0, y1 = (view["years"][0], str(view["years"][-1])[-2:]) if view["years"] else ("", "")
    txt(slide, "Rectangle 2", f"Year-wise themes connect execution to the {y0}-{y1} runway")
    tbl = sh(slide, "Table 13")
    themes = view["themes"]
    if themes:
        idx = plan["themes"]["groups"][group_index]
        if view.get("delta"):
            # dedicated "Δ vs previous" column, same treatment as the enabling table
            tm = theme_delta_map(view)
            vd.set_table_columns(tbl, THEME_DELTA_WIDTHS)
            rows, tones = [], []
            for i in idx:
                c = list(themes[i]["cells"])
                text, tone = theme_delta_mark(tm.get((str(c[0]), c[1])))
                rows.append(c[:3] + [marked(text, tone), c[3]])
                tones.append(tone_palette(tone))
            vd.fill_table(tbl, rows, status_col=4, status_values=[themes[i]["status"] for i in idx],
                          header=["Year", "Key theme", "Actions", "Δ progress vs previous", "Progress / status"],
                          palette=STATUS2)
            vd.colour_column(tbl, 3, tones)
            vd.set_body_font(tbl, 9)
            extra_added.append("themes table: added the 'delta progress vs previous' column")
        else:
            rows = [themes[i]["cells"] for i in idx]
            vd.fill_table(tbl, rows, status_col=3, status_values=[themes[i]["status"] for i in idx], palette=STATUS2)
        p = plan["themes"]
        if p["pt"] and p["orig_pt"] and p["pt"] != p["orig_pt"] and not view.get("delta"):
            vd.set_body_font(tbl, p["pt"])
            if group_index == 0:
                extra_added.append(f"body font reduced from {p['orig_pt']:g}pt to {p['pt']:g}pt on the themes table to keep it on one slide")
    else:
        vd.fill_table(tbl, [["", "", "No key themes recorded in the sheet", ""]])
        extra_added.append("placeholder text 'No key themes recorded in the sheet' on the themes table")


def build_dependencies(slide, view, model, rows_idx, extra_added, source_text):
    deps = view["dep_register"]
    q = view["quarter"]
    all_deps = model["dependencies"]
    n = len(all_deps)
    n_ns = sum(1 for d in all_deps if d["status"] == "Not started")
    functions = sorted({d["function"] for d in all_deps}, key=str.lower)
    ly = view["last_year"]
    txt(slide, "Rectangle 7", str(n))
    txt(slide, "Rectangle 9", f"Across {len(functions)} function{'s' if len(functions) != 1 else ''}" if n else "None recorded in the sheet")
    txt(slide, "Rectangle 12", str(n_ns))
    txt(slide, "Rectangle 14", f"{round(100 * n_ns / n)}% of dependency actions" if n else "—")
    if n_ns and n:
        by_f = {}
        for d in all_deps:
            if d["status"] == "Not started":
                by_f[d["function"]] = by_f.get(d["function"], 0) + 1
        top = sorted(by_f.items(), key=lambda x: -x[1])[:2]
        if len(top) == 2:
            focus = f"{top[0][0]} and {top[1][0]} hold {top[0][1] + top[1][1]} of the {n_ns} not-started actions."
        else:
            focus = f"{top[0][0]} holds all {n_ns} not-started actions."
    elif n:
        focus = f"All {n} dependency actions have started."
    else:
        focus = "No interdependency actions are recorded for this function."
    txt(slide, "Rectangle 22", focus)
    n_last = sum(1 for d in all_deps if d["year"] == ly)
    txt(slide, "Rectangle 26", f"{ly}: no actions captured" if not n_last else f"{ly}: {n_last} action{'s' if n_last != 1 else ''} captured")
    if n_last:
        fill(sh(slide, "Rectangle: Rounded Corners 25"), C["blue_bg"])
        color(sh(slide, "Rectangle 26"), C["blue"])
    txt(slide, "Rectangle 121", source_text)
    # month-on-month status delta above the register (how many were in progress last month vs now)
    if view.get("delta"):
        ds = view["delta"]["dep_status"]
        items = []
        for label, tone_up in (("Completed", "up"), ("In progress", "new"), ("Not started", "down")):
            pv, cv = ds["prev"].get(label, 0), ds["cur"].get(label, 0)
            diff = cv - pv
            tone = "flat" if diff == 0 else (tone_up if diff > 0 else ("down" if tone_up == "up" else "up"))
            items.append((f"{label} {pv} → {cv} ({diff:+d})", tone))
        w, gap, x = 1.75, 0.12, 4.60
        for i, (text, tone) in enumerate(items):
            chip(slide, x + i * (w + gap), 2.78, w, 0.30, text, tone, size=7.0)
        extra_added.append("dependency register: month-on-month status chips added above the register")
    # register rows: clear template example rows, rebuild from prototypes
    proto = {n_: sh(slide, n_) for n_ in ("Rectangle 32", "Rectangle 33", "Rectangle 34", "Rectangle 35", "Rectangle: Rounded Corners 36", "Rectangle 37")}
    keep = set(proto)
    for s in shapes_in_band(slide, 3.44, 6.60, exclude=keep):
        remove(s)
    rows = [deps[i] for i in rows_idx]
    if not rows:
        rows = [{"year": "", "function": "", "action": "No interdependency actions recorded in the sheet", "status": None}]
        extra_added.append("placeholder text 'No interdependency actions recorded in the sheet' on the dependency register")
    for i, d in enumerate(rows):
        dy = i * REGISTER_PITCH
        bg = proto["Rectangle 32"] if i == 0 else (clone(slide, proto["Rectangle 32"], dy=dy) if i % 2 == 0 else None)
        y_ = proto["Rectangle 33"] if i == 0 else clone(slide, proto["Rectangle 33"], dy=dy)
        f_ = proto["Rectangle 34"] if i == 0 else clone(slide, proto["Rectangle 34"], dy=dy)
        a_ = proto["Rectangle 35"] if i == 0 else clone(slide, proto["Rectangle 35"], dy=dy)
        pb = proto["Rectangle: Rounded Corners 36"] if i == 0 else clone(slide, proto["Rectangle: Rounded Corners 36"], dy=dy)
        pt_ = proto["Rectangle 37"] if i == 0 else clone(slide, proto["Rectangle 37"], dy=dy)
        vd.set_shape_text(y_, str(d["year"]))
        vd.set_shape_text(f_, d["function"] or "")
        vd.set_shape_text(a_, d["action"] or vd.NOT_AVAILABLE)
        st = d["status"]
        pal = STATUS2.get(st, STATUS2["No data"])
        fb, fc = pal["fill"], pal["font"]
        vd.set_shape_text(pt_, st or "—")
        fill(pb, fb)
        color(pt_, fc)
    # z-order is naturally correct: each row's background is appended before that row's texts and rows do not overlap


def build_function_progress(slide, view, model, extra_added, note):
    """Which supporting function progressed: average dependency progress, previous vs current."""
    d = view["delta"]
    fns = d["dep_functions"]
    moved = [f for f in fns if f["delta"] not in (None, 0)]
    flat = [f for f in fns if f["delta"] == 0]
    newf = [f for f in fns if f["n_prev"] == 0 and f["n_cur"]]
    txt(slide, "Rectangle 1", "SUPPORT PROGRESS BY FUNCTION")
    txt(slide, "Rectangle 2", "Which supporting function progressed")
    txt(slide, "Rectangle 3", f"Average progress of each function's cross-functional actions, {d['previous_file']} versus {d['current_file']}. "
                              f"Bars are the average of the progress values recorded in the tracker for that function's actions.")
    txt(slide, "Rectangle 12", str(len(fns)))
    txt(slide, "Rectangle 13", "Supporting functions")
    fill(sh(slide, "Rectangle 11"), C["blue"])
    color(sh(slide, "Rectangle 12"), C["blue"])
    txt(slide, "Rectangle 52", str(len(moved)))
    txt(slide, "Rectangle 54", "Functions that moved")
    txt(slide, "Rectangle 56", f"{len(flat)} unchanged, {len(newf)} new in this tracker")
    if not moved:
        fill(sh(slide, "Rectangle 50"), C["amber"])
        color(sh(slide, "Rectangle 52"), C["amber"])
    txt(slide, "Rectangle 16", "Average action progress by supporting function")
    txt(slide, "Rectangle 18", "Previous")
    txt(slide, "Rectangle 20", "Current")
    for n in ("Rectangle: Rounded Corners 21", "Rectangle 22"):
        remove(sh(slide, n))
    fill(sh(slide, "Rectangle: Rounded Corners 17"), "C9CFD6")
    fill(sh(slide, "Rectangle: Rounded Corners 19"), C["blue"])
    old = sh(slide, "Chart")                     # a copied chart shares the template's chart part
    left, top, width, height = old.left, old.top, old.width, old.height
    remove(old)
    cats = [f["function"] for f in fns] or ["No dependency rows"]
    prev = [f["prev"] if f["prev"] is not None else 0 for f in fns] or [0]
    cur = [f["cur"] if f["cur"] is not None else 0 for f in fns] or [0]
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Previous", prev)
    cd.add_series("Current", cur)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, cd)
    gf.name = "Chart Function Progress"
    ch = gf.chart
    ch.has_title = False
    ch.font.name, ch.font.size = "Arial", Pt(8)
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 72, -8
    for si, (ser, colr, vals) in enumerate(((plot.series[0], "C9CFD6", prev), (plot.series[1], C["blue"], cur))):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(colr)
        for pi, v in enumerate(vals):
            dl = ser.points[pi].data_label
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.text_frame.text = f"{v}%"
            r_ = dl.text_frame.paragraphs[0].runs[0]
            r_.font.size, r_.font.bold = Pt(7), True
            r_.font.color.rgb = rgb(C["muted"] if si == 0 else C["blue"])
    cat = ch.category_axis
    cat.tick_labels.font.size = Pt(8)
    cat.has_major_gridlines = False
    cat.format.line.color.rgb = rgb(C["line"])
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale, va.major_unit = 100, 0, 25
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(C["line"])
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0"%"', False
    va.tick_labels.font.size = Pt(7)
    va.format.line.fill.background()
    # per-function read-out in the remarks panel
    txt(slide, "Rectangle 26", "Function read-out")
    proto = {n: sh(slide, n) for n in ("Rectangle 27", "Rectangle 28", "Rectangle 29")}
    for n in ("Rectangle 30", "Rectangle 31", "Rectangle 32", "Rectangle 33", "Rectangle 34"):
        remove(sh(slide, n))
    n = max(len(fns), 1)
    top0, pitch = 2.27, min(1.47, (6.69 - 2.27) / n)
    for i, f in enumerate(fns or [None]):
        lab = proto["Rectangle 27"] if i == 0 else clone(slide, proto["Rectangle 27"], dy=i * pitch)
        body = proto["Rectangle 28"] if i == 0 else clone(slide, proto["Rectangle 28"], dy=i * pitch)
        if f is None:
            vd.set_shape_text(lab, "NO DATA")
            vd.set_shape_text(body, "The tracker records no cross-functional dependency actions for this function.")
            break
        vd.set_shape_text(lab, f["function"].upper())
        if f["n_prev"] == 0:
            line, col = f"New in this tracker: {f['n_cur']} actions at {f['cur']}% average progress.", C["blue"]
        elif f["delta"] == 0:
            line, col = f"No change: {f['n_cur']} actions, {f['cur']}% average progress (was {f['prev']}%).", C["muted"]
        else:
            line = f"{f['delta']:+d} pts: {f['prev']}% → {f['cur']}% across {f['n_cur']} actions."
            col = C["green"] if f["delta"] > 0 else C["red"]
        color(lab, col)
        vd.set_shape_text(body, line)
        body.height = Emu(int(Inches(pitch - 0.16)))
        if i < len(fns) - 1:
            s_ = proto["Rectangle 29"] if i == 0 else clone(slide, proto["Rectangle 29"], dy=i * pitch)
            s_.top = Emu(int(Inches(top0 + (i + 1) * pitch - 0.08)))
        elif i == 0:
            remove(proto["Rectangle 29"])
    vd.set_notes(slide, note)
    extra_added.append("added the 'Which supporting function progressed' slide (previous vs current average action progress)")


def build_impact(slide, view, model, crit_states, rows_idx, extra_added):
    q = view["quarter"]
    fn = model["function"]
    outs = view.get("outcomes") or []
    years = view["years"]
    y0, y1 = (years[0], years[-1]) if years else (None, None)
    crit = model["kpis"]["critical"]
    names = [o["metric"] for o in outs]
    txt(slide, "Rectangle 2", f"How {fn} KPIs translate into VCP value")
    if outs:
        txt(slide, "Rectangle 3", f"The workbook quantifies the {', '.join(names)} trajector{'y' if len(names) == 1 else 'ies'} for {y0}-{y1}; the driver map links each critical KPI to them.")
    else:
        txt(slide, "Rectangle 3", "The workbook has no VCP outcome table for this function; outcome boxes are marked 'Not quantified'.")

    def fy(o, y):
        v = o["years"].get(str(y))
        return vd.fmt_value(v, "currency") if isinstance(v, (int, float)) else (str(v) if v is not None else None)
    # cards: first three outcome metrics (or 'Not quantified')
    card_ids = (("Rectangle 7", "Rectangle 8", "Rectangle 9"), ("Rectangle 12", "Rectangle 13", "Rectangle 14"), ("Rectangle 17", "Rectangle 18", "Rectangle 19"))
    for i, (big, lab, sub) in enumerate(card_ids):
        if i < len(outs):
            o = outs[i]
            txt(slide, big, fy(o, y1) or vd.NOT_AVAILABLE)
            txt(slide, lab, f"{y1} VCP {o['metric']} trajectory")
            txt(slide, sub, f"From {fy(o, y0) or vd.NOT_AVAILABLE} in {y0}")
        else:
            txt(slide, big, "Not quantified")
            txt(slide, lab, "VCP outcome" if i else "VCP outcome trajectory")
            txt(slide, sub, "No outcome table in the sheet for this metric")
    # driver map: one row per critical KPI (template rows: labels 22/34/46; boxes 23-33, 35-45, 47-57)
    row_names = [("Rectangle 22", "Rectangle: Rounded Corners 23", "Rectangle 24", "Rectangle 25", "Arrow: Chevron 26", "Rectangle: Rounded Corners 27", "Rectangle 28", "Rectangle 29", "Arrow: Chevron 30", "Rectangle: Rounded Corners 31", "Rectangle 32", "Rectangle 33"),
                 ("Rectangle 34", "Rectangle: Rounded Corners 35", "Rectangle 36", "Rectangle 37", "Arrow: Chevron 38", "Rectangle: Rounded Corners 39", "Rectangle 40", "Rectangle 41", "Arrow: Chevron 42", "Rectangle: Rounded Corners 43", "Rectangle 44", "Rectangle 45"),
                 ("Rectangle 46", "Rectangle: Rounded Corners 47", "Rectangle 48", "Rectangle 49", "Arrow: Chevron 50", "Rectangle: Rounded Corners 51", "Rectangle 52", "Rectangle 53", "Arrow: Chevron 54", "Rectangle: Rounded Corners 55", "Rectangle 56", "Rectangle 57")]
    proto = [sh(slide, n) for n in row_names[0]]
    for names_ in row_names[1:]:
        for n in names_:
            remove(sh(slide, n))
    kpis = [(crit[i], crit_states[i]) for i in rows_idx] if rows_idx else []
    n = max(len(kpis), 1)
    pitch = 1.08 if n <= 3 else (6.69 - 3.21) / n
    scale = min(1.0, (pitch - 0.41) / 0.67)   # box height shrink when 4 rows
    for i, (k, st) in enumerate(kpis or [(None, None)]):
        row = proto if i == 0 else [clone(slide, p, dy=i * pitch) for p in proto]
        lab, b1, b1t, b1x, _, b2, b2t, b2x, _, b3, b3t, b3x = row
        if scale < 1.0:
            for b in (b1, b2, b3):
                b.height = Emu(int(b.height * scale))
        if k is None:
            vd.set_shape_text(lab, "NO CRITICAL KPI")
            vd.set_shape_text(b1t, "Committed levers"); vd.set_shape_text(b1x, "No critical KPI rows in the sheet")
            vd.set_shape_text(b2t, "Critical KPI"); vd.set_shape_text(b2x, vd.NOT_AVAILABLE)
            vd.set_shape_text(b3t, "VCP outcome"); vd.set_shape_text(b3x, "Not quantified in the sheet")
            continue
        sig, ratio, ft, fa, gap = st
        vd.set_shape_text(lab, k["metric"].upper())
        lab.width = Emu(int(Inches(4.5)))          # template label box is sized for a short example word
        vd.set_shape_text(b1t, "Committed levers")
        vd.set_shape_text(b1x, k["committed"] or "Committed action not captured in the sheet")
        fill(b1, C["blue_bg"]); color(b1t, C["blue"])
        vd.set_shape_text(b2t, k["metric"])
        if ratio is not None:
            vd.set_shape_text(b2x, f"{round(100 * ratio)}% of {q} target; gap {gap_abs(k, gap)}" if sig == "Below target" else f"{round(100 * ratio)}% of {q} target ({fa} vs {ft})")
        else:
            vd.set_shape_text(b2x, f"No {q} target/actual pair in the sheet")
        bg, fc = (C["red_bg"], C["red"]) if sig == "Below target" else ((C["green_bg"], C["green"]) if sig == "On / above target" else (C["grey_bg"], C["muted"]))
        fill(b2, bg); color(b2t, fc)
        if i < len(outs):
            o = outs[i]
            vd.set_shape_text(b3t, f"VCP {o['metric']}")
            vd.set_shape_text(b3x, f"{fy(o, y0) or vd.NOT_AVAILABLE} in {y0} to {fy(o, y1) or vd.NOT_AVAILABLE} in {y1}")
            fill(b3, C["blue_bg"]); color(b3t, C["blue"])
        else:
            vd.set_shape_text(b3t, "VCP outcome")
            vd.set_shape_text(b3x, "Not quantified in the sheet")
            fill(b3, C["amber_bg"]); color(b3t, C["amber"])
    # trajectory chart + chip
    gf = sh(slide, "Chart")
    if outs and years:
        cats = [str(y) for y in years]
        series = []
        labels = []
        for o in outs:
            vals = [o["years"].get(str(y)) if isinstance(o["years"].get(str(y)), (int, float)) else 0 for y in years]
            series.append((o["metric"], vals))
            labels.append([fy(o, y) or "" for y in years])
        update_chart(gf, cats, series, series_colours=SERIES_COLOURS[:len(series)], labels=labels, label_pt=7 if len(series) <= 2 else 6)
        chips = [f"{y1} {o['metric']} plan: {vd.fmt_value(o['pct'][str(y1)], 'percent')}" for o in outs if o.get("pct", {}).get(str(y1)) is not None]
        txt(slide, "Rectangle 62", "; ".join(chips) if chips else f"Values in €M from the sheet's outcome table ({y0}-{y1})")
        txt(slide, "Rectangle 59", "VCP outcome trajectory (€M)")
    else:
        update_chart(gf, [str(y) for y in years] or ["—"], [("Not quantified", [0] * (len(years) or 1))], series_colours=[C["blue"]])
        txt(slide, "Rectangle 62", "No outcome table in the sheet")
        txt(slide, "Rectangle 59", "VCP outcome trajectory (€M) — not quantified")


def build_rca(slide, view, model, crit_states, rows_idx, extra_added):
    q = view["quarter"]
    rr = view["readiness_raw"]
    reds = view["rca"]["rows"]
    all_kpis = model["kpis"]["critical"] + model["kpis"]["enabling"]
    by_name = {k["metric"]: k for k in all_kpis}
    txt(slide, "Rectangle 2", f"{q} red KPIs need documented root causes" if reds else f"No {q} red KPIs; data gaps still limit the review")
    txt(slide, "Rectangle 3", f"RCA is not captured in the sheet; {rr['q_missing']} of {rr['q_total']} quarterly target/actual cells are blank.")
    txt(slide, "Rectangle 6", f"RCA follow-up for {q} red KPIs")
    proto = [sh(slide, n) for n in ("Rectangle: Rounded Corners 7", "Rectangle 8", "Rectangle 9", "Rectangle: Rounded Corners 10", "Rectangle 11", "Rectangle 12", "Rectangle 13")]
    for names_ in (("Rectangle: Rounded Corners 14", "Rectangle 15", "Rectangle 16", "Rectangle: Rounded Corners 17", "Rectangle 18", "Rectangle 19", "Rectangle 20"),
                   ("Rectangle: Rounded Corners 21", "Rectangle 22", "Rectangle 23", "Rectangle: Rounded Corners 24", "Rectangle 25", "Rectangle 26", "Rectangle 27"),
                   ("Rectangle: Rounded Corners 28", "Rectangle 29", "Rectangle 30", "Rectangle: Rounded Corners 31", "Rectangle 32", "Rectangle 33", "Rectangle 34")):
        for n in names_:
            remove(sh(slide, n))
    items = [reds[i] for i in rows_idx] if rows_idx else []
    n = max(len(items), 1)
    pitch = 1.04 if n <= 4 else (6.60 - 2.14) / n
    for i, r in enumerate(items or [None]):
        row = proto if i == 0 else [clone(slide, p, dy=i * pitch) for p in proto]
        card, bar, name, pillbg, pill, l1, l2 = row
        if pitch < 1.04:
            card.height = Emu(int(Inches(pitch - 0.09))); bar.height = card.height
        if r is None:
            vd.set_shape_text(name, f"No KPI below the {q} target")
            vd.set_shape_text(pill, "—")
            vd.set_shape_text(l1, f"{rr['ready_total']} KPIs have a {q} target/actual pair; none is below target.")
            vd.set_shape_text(l2, "RCA status: nothing to document this period.")
            fill(bar, C["green"]); color(name, C["green"]); fill(pillbg, C["green_bg"]); color(pill, C["green"])
            continue
        metric, signal_txt = r[0], r[1]
        k = by_name.get(metric)
        gap = vd.delta_text(k, q).strip("()") if k else "—"
        vd.set_shape_text(name, metric)
        vd.set_shape_text(pill, gap)
        vd.set_shape_text(l1, f"RCA status: not captured in the sheet. {q} actual vs target: {signal_txt}.")
        vd.set_shape_text(l2, f"Committed action (sheet): {(k or {}).get('committed') or 'not captured in the sheet'}")
    # sheet performance report (all counts computed from the sheet)
    txt(slide, "Rectangle 38", f"{rr['q_missing']}/{rr['q_total']}")
    pop = (rr["q_total"] - rr["q_missing"]) / rr["q_total"] if rr["q_total"] else 0
    bar = sh(slide, "Rectangle: Rounded Corners 41")
    bar.width = Emu(int(max(Inches(0.02), Inches(3.12) * pop)))
    txt(slide, "Rectangle 42", f"{round(100 * pop)}% populated")
    txt(slide, "Rectangle 45", f"{rr['crit_pop_pct']}%")
    txt(slide, "Rectangle 50", f"{rr['enab_pop_pct']}%")
    fills = [("Rectangle 54", "Rectangle: Rounded Corners 56", "Rectangle 57", "Owner", rr["owner_blank"]),
             ("Rectangle 58", "Rectangle: Rounded Corners 60", "Rectangle 61", "Baseline (2025)", rr["baseline_missing"])]
    yrs = view["years"] + [None, None, None]
    for j, (labn, barn, cntn) in enumerate((("Rectangle 62", "Rectangle: Rounded Corners 64", "Rectangle 65"), ("Rectangle 66", "Rectangle: Rounded Corners 68", "Rectangle 69"), ("Rectangle 70", "Rectangle: Rounded Corners 72", "Rectangle 73"))):
        y = yrs[j]
        fills.append((labn, barn, cntn, f"Target {y}" if y else "Target —", rr["targets_missing"].get(str(y), 0) if y else 0))
    fills.append(("Rectangle 74", "Rectangle: Rounded Corners 76", "Rectangle 77", "Committed actions", rr["committed_missing"]))
    nk = rr["n_kpis"] or 1
    for labn, barn, cntn, label, missing in fills:
        txt(slide, labn, label)
        txt(slide, cntn, f"{missing}/{rr['n_kpis']}")
        frac = missing / nk
        b = sh(slide, barn)
        b.width = Emu(int(max(Inches(0.02), Inches(1.96) * frac)))
        colr = C["red"] if frac >= 0.999 else (C["amber"] if frac > 0 else C["green"])
        fill(b, colr)
        color(sh(slide, cntn), colr)
    empty = rr["theme_cols_empty"]
    txt(slide, "Rectangle 79", f"{len(empty)} fully empty theme-tracker column{'s' if len(empty) != 1 else ''}")
    na = rr["n_actions"]
    txt(slide, "Rectangle 80", (", ".join(empty) + ". " if empty else "No fully empty theme-tracker column. ")
        + f"Planned Completion is blank in {rr['planned_blank']}/{na} rows; Action Progress in {rr['progress_blank']}/{na}.")


# ----------------------------------------------------------------------------- delta slides (v2 style)
def add_delta_slides_v2(prs, view, extra_added, note, sa, sb):
    """sa = clean copy of the critical-KPI slide, sb = clean copy of the themes slide (taken before any filling)."""
    import vcp_delta
    d = view["delta"]
    ts = d["theme_summary"]
    txt(sa, "Rectangle 1", "MONTH-ON-MONTH")
    txt(sa, "Rectangle 2", "Key theme movement vs previous tracker")
    txt(sa, "Rectangle 3", f"Previous: {d['previous_file']}  →  current: {d['current_file']}. Bars show key-theme progress (Key Theme Progress column, else average of action progress).")
    txt(sa, "Rectangle 12", str(ts["compared"])); txt(sa, "Rectangle 13", "Key themes compared")
    fill(sh(sa, "Rectangle 11"), C["blue"]); color(sh(sa, "Rectangle 12"), C["blue"])
    txt(sa, "Rectangle 52", str(ts["moved"] + ts["new"])); txt(sa, "Rectangle 54", "Themes moved forward or newly added")
    txt(sa, "Rectangle 56", f"{ts['unchanged']} unchanged, {ts['declined']} declined, {ts['removed']} removed")
    fill(sh(sa, "Rectangle 50"), C["green"]); color(sh(sa, "Rectangle 52"), C["green"])
    for n in ("Rectangle: Rounded Corners 25", "Rectangle 26", "Rectangle 27", "Rectangle 28", "Rectangle 29", "Rectangle 30", "Rectangle 31", "Rectangle 32", "Rectangle 33", "Rectangle 34",
              "Rectangle: Rounded Corners 17", "Rectangle 18", "Rectangle: Rounded Corners 19", "Rectangle 20", "Rectangle: Rounded Corners 21", "Rectangle 22"):
        try:
            remove(sh(sa, n))
        except KeyError:
            pass
    card = sh(sa, "Rectangle: Rounded Corners 15")
    card.width = Emu(int(Inches(11.94)))
    txt(sa, "Rectangle 16", "Progress by key theme: previous (grey) vs current (blue)")
    old = sh(sa, "Chart")
    remove(old)
    themes = d["themes"]
    cd = CategoryChartData()
    cd.categories = [f"{t['year']} · {vcp_delta._short(t['theme'])}" for t in themes] or ["No key themes"]
    cd.add_series("Previous", [t["prev_pct"] if t["prev_pct"] is not None else 0 for t in themes] or [0])
    cd.add_series("Current", [t["cur_pct"] if t["cur_pct"] is not None else 0 for t in themes] or [0])
    gf = sa.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.81), Inches(3.4), Inches(11.4), Inches(3.65), cd)
    gf.name = "Chart Delta"
    ch = gf.chart
    ch.has_title = False
    ch.font.name = "Arial"; ch.font.size = Pt(8)
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(8)
    plot = ch.plots[0]; plot.gap_width = 55; plot.overlap = -12; plot.has_data_labels = True
    dl = plot.data_labels; dl.number_format = '0"%"'; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(7)
    plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb = rgb("C9CFD6")
    plot.series[1].format.fill.solid(); plot.series[1].format.fill.fore_color.rgb = rgb(C["blue"])
    cat = ch.category_axis; cat.reverse_order = True; cat.tick_labels.font.size = Pt(8); cat.has_major_gridlines = False; cat.format.line.color.rgb = rgb(C["line"])
    val = ch.value_axis; val.minimum_scale, val.maximum_scale, val.major_unit = 0, 100, 25
    val.has_major_gridlines = True; val.major_gridlines.format.line.color.rgb = rgb(C["line"])
    val.tick_labels.number_format = '0"%"'; val.tick_labels.number_format_is_linked = False; val.tick_labels.font.size = Pt(7)
    val.tick_label_position = XL_TICK_LABEL_POSITION.LOW; val.format.line.fill.background()
    vd.set_notes(sa, note)
    if sb is None:
        return
    # change table on the clean copy of the themes slide
    txt(sb, "Rectangle 1", "MONTH-ON-MONTH")
    title = "What changed since the previous tracker"
    txt(sb, "Rectangle 2", title)
    rows = [[c["year"], c["theme"], c["change"], c["movement"]] for c in d["changes"]]
    statuses = [TONE2.get(c["tone"]) for c in d["changes"]]
    if not rows:
        rows = [["", "", f"No changes between {d['previous_file']} and {d['current_file']}", ""]]
        statuses = [None]
        extra_added.append("placeholder text 'No changes between the two trackers' on the delta table (no changes found)")
    header = ["Year", "Key theme / KPI", "Change", "Previous → Current"]
    geo = vd._table_geometry(sh(sb, "Table 13"))
    groups = vd.pack_rows(rows, geo)
    vd.fill_section(prs, sb, "Table 13", rows, statuses, 3, header, groups, title, extra_added, note=note, palette=STATUS2, title_shape="Rectangle 2")
    extra_added.append("month-on-month delta slides added on request: 'Key theme movement vs previous tracker' (native chart) and 'What changed since the previous tracker'")


# ----------------------------------------------------------------------------- main entry
def write_deck_v2(view, model, template, out_path, plan, allow_extra, impact_mode, extra_added, note_for):
    prs = vd.Presentation(str(template))
    slides = list(prs.slides)
    s1, s2, s3, s4, s5, s6, s7 = slides[:7]
    q = view["quarter"]
    src = view["sources"]

    def dup_group(base, n):
        """base + (n-1) copies inserted right after it, made BEFORE any filling (clean template copies)."""
        out = [base]
        for i in range(1, n):
            out.append(vd.duplicate_slide(prs, base, prs.slides.index(base) + i - 1))
        return out

    n_dep_slides = max(1, len(plan["dependencies"]["groups"]))
    n_rca_slides = max(1, len(plan["rca"]["groups"]))
    n_enab = max(1, len(plan["enabling"]["groups"]))
    n_them = max(1, len(plan["themes"]["groups"]))
    drop_impact = impact_mode == "drop" and not view.get("outcomes")
    n_imp = 0 if drop_impact else max(1, len(plan["impact"]["groups"]))
    # duplicates first (copies of untouched template slides)
    enab_slides = dup_group(s3, n_enab)
    them_slides = dup_group(s4, n_them)
    dep_slides = dup_group(s5, n_dep_slides)
    imp_slides = dup_group(s6, n_imp) if n_imp else []
    rca_slides = dup_group(s7, n_rca_slides)
    # ---- KPI cockpit / linkage + one RCA drill-down per red KPI: part of every deck.
    # Governance decision slides are opt-in (--with-governance-slides).
    end = lambda: len(prs.slides._sldIdLst) - 1
    q_ = view["quarter"]
    all_k = model["kpis"]["critical"] + model["kpis"]["enabling"]
    red_k = [k for k in all_k if vd.signal_for(k, q_, [], model["sheet"])[0] == "Below target"]
    gov_slides = None
    if view.get("governance") and view.get("with_governance_slides"):
        gp = view["governance_plan"]
        gov_slides = {"decisions": vd.duplicate_slide(prs, s4, end()) if view.get("decisions") else None,
                      "actions": [vd.duplicate_slide(prs, s3, end()) for _ in gp["action_parts"]]}
    kpi_slides = {"linkage": vd.duplicate_slide(prs, s3, end()),
                  "rca": [(k, vd.duplicate_slide(prs, s5, end())) for k in red_k]}
    fnprog_slide = None
    if view.get("delta"):        # always present, so every function gets the same slide set
        fnprog_slide = vd.duplicate_slide(prs, s2, prs.slides.index(dep_slides[-1]))
    delta_slides = None
    if view.get("delta"):                     # clean copies taken before any filling, placed after the last themes slide
        sa = vd.duplicate_slide(prs, s2, prs.slides.index(them_slides[-1]))
        sb = vd.duplicate_slide(prs, s4, prs.slides.index(sa)) if view.get("with_change_table") else None
        delta_slides = (sa, sb)
    if drop_impact:
        vd.delete_slide(prs, s6)
        extra_added.append("'VCP impact linkage' slide dropped on request (no outcome table in the sheet)")

    # cover
    vd.set_shape_text(sh(s1, "Text Placeholder 10"), view["title"])
    vd.set_notes(s1, note_for("cover"))
    # slide 2
    txt(s2, "Rectangle 1", f"{model['function'].upper()} · CRITICAL VCP COMMITMENTS")
    crit_states = build_critical(s2, view, model, extra_added)
    crit_states = [kpi_state(k, q, [], model["sheet"]) for k in model["kpis"]["critical"]]   # full tuples for later slides
    vd.set_notes(s2, note_for("critical"))
    # slide 3(+)
    for gi, s in enumerate(enab_slides):
        build_enabling(s, view, model, plan, extra_added, gi)
        if gi:
            txt(s, "Rectangle 2", f"Enabling KPIs (cont.)")
            extra_added.append("VCP Commitment Enabling KPIs (cont.)")
        vd.set_notes(s, note_for("enabling"))
    # slide 4(+)
    for gi, s in enumerate(them_slides):
        build_themes(s, view, plan, extra_added, gi)
        if gi:
            txt(s, "Rectangle 2", "Year-wise themes (cont.)")
            extra_added.append("Themes & Action Items (cont.)")
        vd.set_notes(s, note_for("themes"))
    # delta slides (after the themes slides)
    if delta_slides:
        add_delta_slides_v2(prs, view, extra_added, note_for("delta"), *delta_slides)
    # slide 5(+)
    groups = plan["dependencies"]["groups"] or [[]]
    for gi, s in enumerate(dep_slides):
        build_dependencies(s, view, model, groups[gi] if gi < len(groups) else [], extra_added, "Source: " + src["dependencies"])
        if gi:
            txt(s, "Rectangle 2", "Cross-functional actions (cont.)")
            extra_added.append("Cross-functional dependencies (cont.)")
        vd.set_notes(s, note_for("dependencies"))
    if fnprog_slide is not None:
        build_function_progress(fnprog_slide, view, model, extra_added, note_for("delta"))
    # slide 6(+)
    groups = plan["impact"]["groups"] or [[]]
    for gi, s in enumerate(imp_slides):
        build_impact(s, view, model, crit_states, groups[gi] if gi < len(groups) else [], extra_added)
        if gi:
            txt(s, "Rectangle 2", f"How {model['function']} KPIs translate into VCP value (cont.)")
            extra_added.append("VCP impact linkage (cont.)")
        vd.set_notes(s, note_for("impact"))
    # slide 7(+)
    groups = plan["rca"]["groups"] or [[]]
    for gi, s in enumerate(rca_slides):
        build_rca(s, view, model, crit_states, groups[gi] if gi < len(groups) else [], extra_added)
        if gi:
            txt(s, "Rectangle 2", f"{q} red KPIs (cont.)")
            extra_added.append("RCA for red KPIs (cont.)")
        vd.set_notes(s, note_for("rca"))
    import vcp_governance
    if gov_slides:
        vcp_governance.fill_governance_slides(prs, gov_slides, kpi_slides, view, view["_model"], view["governance"],
                                              view.get("decisions", []), view["governance_plan"], extra_added, note_for)
    vcp_governance.fill_kpi_slides(prs, kpi_slides, view, view["_model"], extra_added, note_for)
    set_page_numbers(prs)
    view["template_version"] = 2
    view["expected_v2"] = expected_strings(view, model)
    prs.save(str(out_path))
    return extra_added


def expected_strings(view, model):
    """Deck strings the reconciliation must find verbatim for this template."""
    exp = []
    for row in view["enabling"]["rows"]:
        exp += row
    for t in view["themes"]:
        exp += t["cells"]
    for d in view["dep_register"]:
        exp += [str(d["year"]), d["function"] or "", d["action"] or vd.NOT_AVAILABLE, d["status"] or "—"]
    for r in view["rca"]["rows"]:
        exp.append(r[0])
    for k in model["kpis"]["critical"]:
        exp.append(k["metric"])
        if k["committed"]:
            exp.append(k["committed"])
    if view.get("governance"):
        for a in view["governance_plan"]["actions"]:
            exp += [a["action"], a["function"], a["owner"] or "", str(a["due"] or ""), a["status"] or ""]
            if a.get("remarks"):
                exp.append(str(a["remarks"]))
        exp += list(view.get("decisions", []))
        for k in model["kpis"]["critical"] + model["kpis"]["enabling"]:
            exp.append(k["metric"])
    return [e for e in exp if e]
